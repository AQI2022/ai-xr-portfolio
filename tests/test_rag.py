import io
from pathlib import Path

import pytest
from docx import Document

from ai_xr.rag import Retriever, answer_question, chunk_document, parse_document
from ai_xr.providers import ExtractiveProvider
from ai_xr.store import Store


@pytest.fixture
def retriever(tmp_path):
    store = Store(tmp_path/"test.sqlite")
    for name, text in [("anchors.txt", "HoloLens 空间锚点用于固定虚拟内容。跟踪丢失时提示重新扫描。"),
                       ("cpu.txt", "Unity Profiler 分析帧时间，定位 CPU GPU 渲染性能。")]:
        doc, chunks = chunk_document(name, text.encode())
        store.replace_document(doc, chunks)
    return Retriever(store)


def test_chunks_preserve_source_spans():
    text = "空间锚点和坐标变换。"*100
    doc, chunks = chunk_document("../notes.txt", text.encode(), 100, 20)
    assert len(chunks) > 1
    assert all(c["title"] == "notes.txt" and c["text"] == text[c["start"]:c["end"]] for c in chunks)
    assert chunks[-1]["end"] == len(text)
    assert len({c["id"] for c in chunks}) == len(chunks)


def test_document_idempotence(retriever):
    before = len(retriever.store.all_chunks())
    doc, chunks = chunk_document("anchors.txt", "HoloLens 空间锚点用于固定虚拟内容。跟踪丢失时提示重新扫描。".encode())
    retriever.store.replace_document(doc, chunks)
    assert len(retriever.store.all_chunks()) == before


def test_chinese_and_exact_keyword_search(retriever):
    assert retriever.search("空间锚点", 1)[0]["title"] == "anchors.txt"
    assert retriever.search("Profiler GPU", 1)[0]["title"] == "cpu.txt"


def test_no_evidence_returns_empty(retriever):
    assert retriever.search("xylophonicbanana123", 3) == []


def test_docx_table_extraction():
    doc = Document()
    doc.add_paragraph("XR知识")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "空间锚点"
    stream = io.BytesIO()
    doc.save(stream)
    text = parse_document("test.docx", stream.getvalue())[0][1]
    assert "XR知识" in text and "空间锚点" in text


@pytest.mark.parametrize("size,overlap", [(50, 10), (100, 100), (100, -1)])
def test_bad_chunk_parameters(size, overlap):
    with pytest.raises(ValueError):
        chunk_document("test.txt", b"hello", size, overlap)


async def test_extractive_answer_is_explicit(retriever):
    result = await answer_question("空间锚点", retriever, ExtractiveProvider())
    assert result["mode"] == "extractive" and result["citation_valid"]
    assert result["validation"] == "citation_ids_only"


async def test_unknown_model_citation_is_flagged(retriever):
    class BadProvider:
        mode = "test-model"
        async def chat(self, messages):
            return {"content": "空间锚点使用云端服务 [S999]"}
    result = await answer_question("空间锚点", retriever, BadProvider())
    assert not result["generation_citation_valid"]
    assert result["citation_valid"] and result["evidence_fallback"]
    assert "S999" not in result["answer"]


async def test_abstention(retriever):
    result = await answer_question("xylophonicbanana123", retriever, ExtractiveProvider())
    assert result["abstained"] and not result["citations"]
